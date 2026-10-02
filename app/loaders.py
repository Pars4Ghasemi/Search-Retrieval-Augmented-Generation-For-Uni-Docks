"""Read PDF, Word, PowerPoint and text/Markdown files into plain-text sections.

Each section keeps track of where it came from (file name and, when the format
has them, page or slide number) so answers can cite their sources later.
"""

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


@dataclass
class Section:
    text: str
    metadata: dict = field(default_factory=dict)


def _load_pdf(path: Path) -> list[Section]:
    reader = PdfReader(path)
    return [
        Section(text, {"source": path.name, "page": i})
        for i, page in enumerate(reader.pages, start=1)
        if (text := (page.extract_text() or "").strip())
    ]


def _load_docx(path: Path) -> list[Section]:
    doc = DocxDocument(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts).strip()
    return [Section(text, {"source": path.name})] if text else []


def _load_pptx(path: Path) -> list[Section]:
    sections = []
    for i, slide in enumerate(Presentation(path).slides, start=1):
        parts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        text = "\n\n".join(parts).strip()
        if text:
            sections.append(Section(text, {"source": path.name, "slide": i}))
    return sections


def _load_text(path: Path) -> list[Section]:
    text = path.read_text(encoding="utf-8", errors="ignore").strip()
    return [Section(text, {"source": path.name})] if text else []


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".pptx": _load_pptx,
    ".txt": _load_text,
    ".md": _load_text,
}


def load_file(path: Path) -> list[Section]:
    """Load one file. Raises ValueError for unsupported file types."""
    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    return loader(path)


def iter_supported_files(folder: Path):
    """Yield every supported file under `folder`, recursively, in sorted order."""
    for path in sorted(folder.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS and not path.name.startswith("~$"):
            yield path
